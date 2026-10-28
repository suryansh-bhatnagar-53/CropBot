"""Step 16 on the live test database: a sample sheet filled by this test (reviewer "TEST SAMPLE", decisions in a
temporary file, never the committed one), imported, then loaded twice with the snapshot proof, and re-applied
on an emptied database.

  uv run python -m pytest tests/review/test_live_review.py --neo4j-test-db cropbot-test
"""

import datetime as dt
import json
from pathlib import Path

import pytest

from kg.load import facts as lf
from kg.query import FarmContext, answer
from kg.registry import read_entry
from kg.review import importer, sheet
from kg.review.decisions import DECISIONS_CSV, append_decisions, read_decisions
from kg.validate import VALIDATED_DIR
from tests.conftest import reset_and_seed

FIXTURES = Path(__file__).parents[1] / "query" / "fixtures"


def short(rec) -> str:
    return rec["short"][:4] + rec["short"][12:]


@pytest.fixture(scope="module")
def setup(neo4j_test_db, tmp_path_factory):
    if not (VALIDATED_DIR / "tnau_agriculture_2020" / "facts.jsonl").exists():
        pytest.skip("needs data/validated and data/verified for the TNAU guide")
    driver, name = neo4j_test_db
    tmp = tmp_path_factory.mktemp("review")
    decisions = tmp / "decisions.csv"                       # never the committed data/review/expert_decisions.csv
    committed = DECISIONS_CSV.read_bytes() if DECISIONS_CSV.exists() else None
    concepts = reset_and_seed(driver, name)
    entry = read_entry("source:tnau_agriculture_2020")
    built = lf.build(entry, lf.read_inputs(entry, decisions), concepts)
    lf.load(driver, entry, built, name)
    from kg.load.seed import read_seed
    names = {n.id: n.name_en for n in read_seed().nodes if getattr(n, "name_en", None)}
    rows = sheet.rows(entry, built, sheet.current_from_db(driver, name), names)
    path = tmp / "review.xlsx"
    sheet.write_workbook(path, entry, rows, name)
    yield {"driver": driver, "name": name, "entry": entry, "concepts": concepts, "decisions": decisions,
           "path": path, "built": built}
    assert (DECISIONS_CSV.read_bytes() if DECISIONS_CSV.exists() else None) == committed   # untouched


def fill(path: Path, built: dict) -> dict:
    """The sample decisions, by short fact name: grain drying approved, a storage rule rejected, the nursery
    DAP top dressing (scope in unmapped) needs change, and the BPH insecticide list approved."""
    from openpyxl import load_workbook
    by_short = {short(r): r["id"] for r in built["recs"]}
    sample = {"8e1d#2": ("approve", "", "yes"), "f5b0#20": ("reject", "TEST SAMPLE: not a real decision", "no"),
              "b879#19": ("needs change", "TEST SAMPLE: applies only in less fertile nursery soils", "not sure"),
              "e63c#35": ("approve", "", "not sure")}
    wb = load_workbook(path)
    ws = wb[sheet.FACTS]
    keys = [k for k, _, _ in sheet.COLUMNS]
    col = {k: keys.index(k) + 1 for k in keys}
    wanted = {by_short[s]: v for s, v in sample.items()}
    for n in range(2, ws.max_row + 1):
        v = wanted.get(ws.cell(n, col["fact_id"]).value)
        if v:
            for k, x in zip(("decision", "note", "uttarakhand"), v):
                ws.cell(n, col[k], x or None)
            ws.cell(n, col["reviewer"], "TEST SAMPLE")
            ws.cell(n, col["date"], dt.date(2026, 10, 21))
    wb.save(path)
    return wanted


def statuses(driver, name, ids) -> dict:
    found, _, _ = driver.execute_query(
        "MATCH (r:Recommendation) WHERE r.id IN $ids RETURN r.id AS id, r.status AS status, r.reviewed_by AS by, "
        "r.applies_uttarakhand AS utk, r.review_reasons AS reasons", ids=list(ids), database_=name)
    return {r["id"]: r.data() for r in found}


def import_sheet(s) -> dict:
    meta, rows = sheet.read_workbook(s["path"])
    existing = {d["fact_id"]: d for d in read_decisions(s["decisions"])}
    result = importer.check_rows(rows, importer.current_facts(s["built"]), existing, meta["source_id"],
                                 dt.date(2026, 10, 21))
    accepted = importer.stamp(result["accepted"], s["path"].read_bytes())
    importer.apply(s["driver"], accepted, s["name"])
    append_decisions(accepted, s["decisions"])
    return result


def test_import_load_twice_and_rebuild(setup):
    s = setup
    driver, name, entry = s["driver"], s["name"], s["entry"]
    wanted = fill(s["path"], s["built"])
    first = import_sheet(s)
    assert len(first["accepted"]) == 4 and not first["refused"] and len(first["not_reviewed"]) == 109
    expect = {fid: {"approve": "approved", "reject": "rejected", "needs change": "needs_review"}[d]
              for fid, (d, _, _) in wanted.items()}
    got = statuses(driver, name, wanted)
    assert {i: g["status"] for i, g in got.items()} == expect
    assert all(g["by"] == "TEST SAMPLE" for g in got.values())
    change = next(i for i, (d, _, _) in wanted.items() if d == "needs change")
    assert "expert: TEST SAMPLE: applies only in less fertile nursery soils" in got[change]["reasons"]

    before = lf.snapshot(driver, name)
    again = import_sheet(s)                                                   # a re-run is a no-op
    assert again["already"] and not again["accepted"] and lf.snapshot(driver, name) == before

    built = lf.build(entry, lf.read_inputs(entry, s["decisions"]), s["concepts"])
    assert len(built["decisions"]) == 4 and not built["stale_decisions"]
    one = lf.load(driver, entry, built, name)
    two = lf.load(driver, entry, built, name)
    assert one["write"]["created"] == {} or not any(one["write"]["created"].values())
    assert not any(two["write"]["created"].values()) and not two["changed"]
    assert two["snapshot_before"] == two["snapshot_after"] == one["snapshot_after"]
    print("\nimport then load 1: created %s, %s -> %s" % (dict(one["write"]["created"]), before["sha256"][:8],
                                                         one["snapshot_after"]["sha256"][:8]))
    print("load 2: created %s, %s -> %s, %d nodes, %d relationships" % (
        dict(two["write"]["created"]), two["snapshot_before"]["sha256"][:8], two["snapshot_after"]["sha256"][:8],
        two["snapshot_after"]["nodes"], two["snapshot_after"]["relationships"]))
    assert {i: g["status"] for i, g in statuses(driver, name, wanted).items()} == expect

    # an emptied database gets the decisions back from the file on its first load
    reset_and_seed(driver, name)
    lf.load(driver, entry, built, name)
    rebuilt = statuses(driver, name, wanted)
    assert {i: g["status"] for i, g in rebuilt.items()} == expect
    assert "expert: TEST SAMPLE: applies only in less fertile nursery soils" in rebuilt[change]["reasons"]
    again2 = lf.load(driver, entry, built, name)
    assert not any(again2["write"]["created"].values()) and not again2["changed"]
    print("emptied database, load: statuses %s; second load created %s, snapshot %s unchanged" % (
        sorted(g["status"] for g in rebuilt.values()), dict(again2["write"]["created"]),
        again2["snapshot_after"]["sha256"][:8]))

    # the approved grain-drying fact is served without the unverified label
    fx = json.loads((FIXTURES / "q15_grain_drying.json").read_text(encoding="utf-8"))
    out = answer(driver, fx["template"], FarmContext.model_validate(fx["context"]), name)
    served = [f for sl in out["slots"] for f in sl["facts"]]
    approved = [f for f in served if f["citation"]["status"] == "approved"]
    assert approved and all(f["labels"] == [] for f in approved)
