"""Offline tests for the expert review (step 16): the workbook, the import checks and the loader's
re-application of the decisions file. No database; the TNAU workbook test is skipped without data/verified."""

import copy
import datetime as dt

import pytest

from kg.load import facts as lf
from kg.review import importer, sheet
from kg.review.decisions import append_decisions, read_decisions
from kg.validate import VALIDATED_DIR
from tests.load.test_facts import CONCEPTS, ENTRY, inputs, storage

TODAY = dt.date(2026, 10, 21)


def current_of(*facts):
    return importer.current_facts(lf.build(ENTRY, inputs(*facts), CONCEPTS))


def filled(fact_id, h, decision="approve", note="", utk="yes", who="TEST SAMPLE", date="2026-10-21", row=2):
    return {"_row": row, "fact_id": fact_id, "row_hash": h, "decision": decision, "note": note, "uttarakhand": utk,
            "reviewer": who, "date": date}


# --- the workbook -------------------------------------------------------------------------------------

def test_the_guide_names_where_the_source_is_from():
    from kg.registry import read_entry
    assert sheet.guide_lines(read_entry("source:tnau_agriculture_2020")) == sheet.GUIDE_LINES     # unchanged
    gb = sheet.guide_lines(read_entry("source:gbpuat_kharif_advisory_2020"), "Uttarakhand")
    assert "this source is from Uttarakhand." in gb[4] and "Tamil Nadu" not in " ".join(gb)
    assert "national (all States/UTs)" in sheet.guide_lines(read_entry("source:dfpd_uniform_spec_kms_2026_27"))[4]


def test_windows_and_variety_lists_have_a_review_order_scope_and_values():                # schema v1.6
    assert {"VarietyRecommendation", "SowingWindow"} <= set(sheet.ORDER)
    window = {"subtype": "SowingWindow", "window_for": "nursery_sowing", "start_month": 5, "start_day": 16,
              "end_month": 6, "end_day": 7, "raw": {"start_month": "2 nd fortnight May"}, "steps": [],
              "regions": [{"id": "zone:uk_tarai_bhabhar_plain", "origin": "text"}], "crop_ids": ["crop:rice"]}
    assert "window for: nursery sowing" in sheet.scope(window, {})
    assert "window: 16 May to 7 June" in sheet.values(window, {})
    varieties = {"subtype": "VarietyRecommendation", "recommends": ["variety:govind"], "steps": [], "regions": []}
    assert "recommends: Govind" in sheet.values(varieties, {"variety:govind": "Govind"})

def test_a_claim_reads_as_a_plain_statement():
    assert sheet.statement("Does the passage state that, for the crop grown for grain, the produce is dried to "
                           "12% moisture?") == "For the crop grown for grain, the produce is dried to 12% moisture."


@pytest.fixture(scope="module")
def tnau_sheet(tmp_path_factory):
    """The TNAU workbook built offline: statuses and review reasons from step 13's rows stand in for the
    database's."""
    if not (VALIDATED_DIR / "tnau_agriculture_2020" / "facts.jsonl").exists():
        pytest.skip("needs data/validated and data/verified for the TNAU guide")
    from kg.load.seed import read_seed
    from kg.registry import read_entry
    entry = read_entry("source:tnau_agriculture_2020")
    seed = read_seed()
    built = lf.build(entry, lf.read_inputs(entry, tmp_path_factory.mktemp("d") / "none.csv"),
                     {n.id for n in seed.nodes})
    current = {r["id"]: {"status": r["status"], "review_reasons": r["props"]["review_reasons"],
                         "verifier_score": r["props"]["verifier_score"], "missing_in_latest_run": False}
               for r in built["recs"]}
    rows = sheet.rows(entry, built, current, {n.id: n.name_en for n in seed.nodes if getattr(n, "name_en", None)})
    path = tmp_path_factory.mktemp("sheet") / "review.xlsx"
    sheet.write_workbook(path, entry, rows, "offline")
    return path, rows, built


def test_scope_in_unmapped_is_highlighted_on_exactly_the_flagged_facts(tnau_sheet):
    from openpyxl import load_workbook
    path, rows, built = tnau_sheet
    ws = load_workbook(path)[sheet.FACTS]
    col = [k for k, _, _ in sheet.COLUMNS].index("unmapped") + 1
    orange = {ws.cell(n, 2).value for n in range(2, ws.max_row + 1)
              if ws.cell(n, col).fill.fill_type == "solid" and ws.cell(n, col).fill.fgColor.rgb.endswith(sheet.SCOPE_FILL)}
    flagged = {r["id"] for r in built["recs"] if r["fact"].get("scope_in_unmapped")}
    shorts = {r["short"][:4] + r["short"][12:] for r in built["recs"] if r["id"] in orange}   # "b879265e483b#3" -> "b879#3"
    assert orange == flagged and shorts == {"b879#3", "b879#14", "b879#19", "e63c#56"}     # 6 under prompt v3


def test_one_row_per_current_fact_post_harvest_first(tnau_sheet):
    path, rows, built = tnau_sheet
    assert len(rows) == len(built["recs"]) == 113
    order = [sheet.ORDER.index(r["subtype"]) for r in rows]
    assert order == sorted(order) and rows[0]["subtype"] == "DryingSpec" and rows[-1]["subtype"] == "Caution"
    assert all(r["instruction"] and r["quote"] and r["row_hash"].startswith("rh:") for r in rows)


def test_only_the_reviewer_columns_are_editable_and_the_hash_is_hidden(tnau_sheet):
    from openpyxl import load_workbook
    path, rows, _ = tnau_sheet
    wb = load_workbook(path)
    ws = wb[sheet.FACTS]
    keys = [k for k, _, _ in sheet.COLUMNS]
    assert ws.protection.sheet and wb.sheetnames[0] == sheet.GUIDE and wb[sheet.META].sheet_state == "hidden"
    for i, k in enumerate(keys, 1):
        assert ws.cell(2, i).protection.locked == (k not in sheet.REVIEWER), k
    letter = ws.cell(1, keys.index("row_hash") + 1).column_letter
    assert ws.column_dimensions[letter].hidden
    lists = sorted(v.formula1 for v in ws.data_validations.dataValidation if v.type == "list")
    assert lists == ['"approve,reject,needs change"', '"yes,no,not sure"']
    guide = [c.value for c in wb[sheet.GUIDE]["A"] if c.value]
    assert any(g.endswith("until the product's official registration status has been checked.") for g in guide)
    assert any("'for hybrids', 'in BPH prone areas'" in g for g in guide)
    assert guide[-1] == ("• The AI pre-check column is a hint, not a decision: it was made by an AI reading the quote. "
                         "Please check FIX and CHECK rows carefully, and don't rely on OK rows without reading them.")


def test_the_workbook_reads_back_with_its_meta(tnau_sheet):
    path, rows, _ = tnau_sheet
    meta, back = sheet.read_workbook(path)
    assert meta["source_id"] == "source:tnau_agriculture_2020" and len(back) == len(rows)
    assert back[0]["row_hash"] == rows[0]["row_hash"] and back[0]["decision"] is None


# --- the AI pre-check hint -------------------------------------------------------------------------------

def notes_for(rows, results=("FIX", "CHECK")):
    """Synthetic notes: the first rows get `results`, the rest OK."""
    out = []
    for i, r in enumerate(rows):
        res = results[i] if i < len(results) else "OK"
        out.append({"review_no": str(r["review_no"]), "fact_id": r["fact_id"], "ai_precheck": res,
                    "ai_note": "note %d, as written" % i})
    return out


def test_ai_precheck_columns_sit_before_the_reviewer_columns_locked_and_coloured(tnau_sheet, tmp_path):
    from openpyxl import load_workbook
    _, rows, _ = tnau_sheet
    rows = copy.deepcopy(rows)
    hashes = [r["row_hash"] for r in rows]
    assert sheet.add_precheck(rows, notes_for(rows)) == []
    from kg.registry import read_entry
    path = tmp_path / "ai.xlsx"
    sheet.write_workbook(path, read_entry("source:tnau_agriculture_2020"), rows, "offline")
    ws = load_workbook(path)[sheet.FACTS]
    keys = [k for k, _, _ in sheet.COLUMNS]
    assert keys.index("ai_precheck") + 2 == keys.index("decision") and keys.index("ai_note") + 1 == keys.index("decision")
    ai, note, h = (keys.index(k) + 1 for k in ("ai_precheck", "ai_note", "row_hash"))
    assert [ws.cell(n, ai).value for n in (2, 3, 4)] == ["FIX", "CHECK", "OK"] and ws.cell(2, note).value == "note 0, as written"
    for n, colour in ((2, "F8696B"), (3, "FFEB84")):
        assert all(ws.cell(n, c).fill.fgColor.rgb.endswith(colour) for c in (ai, note))
    assert ws.cell(4, ai).fill.fill_type is None                                   # OK: no fill
    assert ws.cell(2, ai).protection.locked and ws.cell(2, note).protection.locked
    assert [ws.cell(n, h).value for n in range(2, len(rows) + 2)] == hashes        # the row hash is unchanged


def test_a_fact_without_a_note_keeps_blank_cells_and_a_note_may_have_no_result(tnau_sheet, tmp_path):
    """2026-10-22: the v1.4 facts the earlier pre-check never saw get no note; an earlier note that no longer
    applies keeps its text with a blank result (no fill)."""
    from openpyxl import load_workbook
    from kg.registry import read_entry
    _, rows, _ = tnau_sheet
    rows = copy.deepcopy(rows)
    ns = notes_for(rows)[:-2]                                   # the last two rows have no note
    ns[2].update(ai_precheck="", ai_note="Earlier AI note no longer applies (fixed in v1.4); not re-checked.")
    assert sheet.add_precheck(rows, ns) == []
    path = tmp_path / "partial.xlsx"
    sheet.write_workbook(path, read_entry("source:tnau_agriculture_2020"), rows, "offline")
    ws = load_workbook(path)[sheet.FACTS]
    keys = [k for k, _, _ in sheet.COLUMNS]
    ai, note = keys.index("ai_precheck") + 1, keys.index("ai_note") + 1
    assert ws.cell(4, ai).value is None and ws.cell(4, note).value.startswith("Earlier AI note no longer applies")
    assert ws.cell(4, ai).fill.fill_type is None
    last = len(rows) + 1
    assert [(ws.cell(n, ai).value, ws.cell(n, note).value) for n in (last - 1, last)] == [(None, None)] * 2


def test_scope_names_the_pests_and_targets(tnau_sheet):
    _, rows, built = tnau_sheet
    by_short = {r["short"][:4] + r["short"][12:]: r["id"] for r in built["recs"]}
    scope = {r["fact_id"]: r["scope"] for r in rows}
    assert "target: Stem borer" in scope[by_short["e63c#57"]] and "target: Leaffolder" in scope[by_short["e63c#58"]]
    assert "target: none (general practice)" in scope[by_short["e63c#2"]]
    assert scope[by_short["e63c#16"]].startswith("pest: Brown planthopper")


@pytest.mark.parametrize("change, problem", [
    (lambda ns: ns[0].update(fact_id="rec:unknown"), "no such fact in the sheet"),
    (lambda ns: ns[1].update(fact_id=ns[0]["fact_id"]), "2 notes"),
    (lambda ns: ns[0].update(ai_precheck="MAYBE"), "is not one of OK, CHECK, FIX"),
    (lambda ns: ns[0].update(ai_precheck="", ai_note=""), "neither a result nor a text"),
])
def test_ai_notes_that_do_not_match_are_refused(tnau_sheet, change, problem):
    _, rows, _ = tnau_sheet
    rows = copy.deepcopy(rows)
    ns = notes_for(rows)
    change(ns)
    with pytest.raises(sheet.SheetError) as e:
        sheet.add_precheck(rows, ns)
    assert any(problem in p for p in e.value.problems)
    assert not any("ai_precheck" in r for r in rows)                                # nothing applied


def test_the_import_ignores_the_ai_columns_and_reads_a_workbook_without_them(tnau_sheet, tmp_path):
    from openpyxl import load_workbook
    from kg.registry import read_entry
    _, rows, built = tnau_sheet
    rows = copy.deepcopy(rows)
    sheet.add_precheck(rows, notes_for(rows))
    path = tmp_path / "ai.xlsx"
    sheet.write_workbook(path, read_entry("source:tnau_agriculture_2020"), rows, "offline")
    _, with_ai = sheet.read_workbook(path)
    assert with_ai[0]["ai_precheck"] == "FIX"
    cur = importer.current_facts(built)
    filled_rows = [{**r, "decision": "approve", "uttarakhand": "yes", "reviewer": "TEST SAMPLE", "date": "2026-10-21"}
                   for r in with_ai[:3]]
    plain = [{k: v for k, v in r.items() if k not in sheet.OPTIONAL} for r in filled_rows]
    assert importer.check_rows(filled_rows, cur, {}, ENTRY.id, TODAY) == importer.check_rows(plain, cur, {}, ENTRY.id, TODAY)
    wb = load_workbook(path)                                         # the first workbook had no AI columns
    ws = wb[sheet.FACTS]
    ws.delete_cols([k for k, _, _ in sheet.COLUMNS].index("ai_precheck") + 1, 2)
    wb.save(tmp_path / "old.xlsx")
    _, old = sheet.read_workbook(tmp_path / "old.xlsx")
    assert len(old) == len(rows) and "ai_precheck" not in old[0] and old[0]["row_hash"] == rows[0]["row_hash"]


# --- the import checks ----------------------------------------------------------------------------------

def test_approve_reject_and_needs_change_are_accepted():
    a, b, c = storage("rec:a"), storage("rec:b", moisture=9.0), storage("rec:c", moisture=10.0)
    cur = current_of(a, b, c)
    rows = [filled("rec:a", cur["rec:a"][0]),
            filled("rec:b", cur["rec:b"][0], "reject", "outdated: no longer advised", "no", row=3),
            filled("rec:c", cur["rec:c"][0], "needs change", "only for long-term seed storage", "not sure", row=4)]
    out = importer.check_rows(rows, cur, {}, ENTRY.id, TODAY)
    assert [d["decision"] for d in out["accepted"]] == ["approve", "reject", "needs change"] and not out["refused"]
    assert importer.report(out)["by_subtype"] == {"StorageRule": {"approve": 1, "reject": 1, "needs change": 1}}


@pytest.mark.parametrize("change, reason", [
    (dict(h="rh:old"), "quote, values or scope changed"),
    (dict(fact_id="rec:gone"), "not a current fact"),
    (dict(decision="maybe"), "is not one of approve"),
    (dict(decision="reject"), "a reject needs a reason"),
    (dict(decision="needs change"), "a needs change needs a reason"),
    (dict(utk=""), "Uttarakhand answer"),
    (dict(who=""), "no reviewer name"),
    (dict(date="21/10/2026"), "is not a date"),
    (dict(date="2027-01-01"), "in the future"),
])
def test_a_bad_row_is_refused_with_its_reason(change, reason):
    a = storage("rec:a")
    cur = current_of(a)
    out = importer.check_rows([filled(**{"fact_id": "rec:a", "h": cur["rec:a"][0], **change})], cur, {}, ENTRY.id,
                              TODAY)
    assert not out["accepted"] and any(reason in r for r in out["refused"][0]["reasons"])


def test_a_changed_fact_is_refused():
    """The sheet was made on the 8 % fact; the book's value is now read as 9 %: same id here, new hash."""
    before = storage("rec:a")
    h = current_of(before)["rec:a"][0]
    after = storage("rec:a", moisture=9.0)
    out = importer.check_rows([filled("rec:a", h)], current_of(after), {}, ENTRY.id, TODAY)
    assert out["refused"][0]["reasons"] == ["the fact's quote, values or scope changed since the sheet was made"]


def test_a_rerun_is_a_no_op_and_a_different_decision_is_refused(tmp_path):
    a = storage("rec:a")
    cur = current_of(a)
    rows = [filled("rec:a", cur["rec:a"][0], "needs change", "add: long term only", "not sure")]
    first = importer.check_rows(rows, cur, {}, ENTRY.id, TODAY)
    path = tmp_path / "decisions.csv"
    append_decisions(importer.stamp(first["accepted"], b"xlsx"), path)
    existing = {d["fact_id"]: d for d in read_decisions(path)}
    again = importer.check_rows(rows, cur, existing, ENTRY.id, TODAY)
    assert again["already"] == ["rec:a"] and not again["accepted"] and not again["refused"]
    other = importer.check_rows([{**rows[0], "decision": "approve"}], cur, existing, ENTRY.id, TODAY)
    assert "already decided" in other["refused"][0]["reasons"][0]


def test_a_row_without_a_decision_is_not_reviewed():
    a = storage("rec:a")
    cur = current_of(a)
    out = importer.check_rows([filled("rec:a", cur["rec:a"][0], decision="")], cur, {}, ENTRY.id, TODAY)
    assert out["not_reviewed"] == ["rec:a"] and not out["accepted"] and not out["refused"]


# --- the loader re-applies the decisions file -------------------------------------------------------------

def decision(fact, decision_="approve", note="", **kw):
    return {"fact_id": fact["id"], "row_hash": lf.row_hash(fact), "decision": decision_, "note": note,
            "applies_uttarakhand": "yes", "reviewed_by": "TEST SAMPLE", "reviewed_at": "2026-10-21",
            "source_id": ENTRY.id, "subtype": fact["subtype"], "quote": "q", **kw}


def test_the_loader_applies_a_matching_decision_and_lists_a_stale_one():
    a, b = storage("rec:a"), storage("rec:b", moisture=9.0)
    changed = copy.deepcopy(b)
    changed["evidence"][0]["quote"] = "Store the seeds for 15 months"                # the quote the expert saw
    ins = inputs(a, b) | {"decisions": [decision(a, "needs change", "long term only"), decision(changed, "reject",
                                                                                                "x"),
                                        decision(storage("rec:gone"))]}
    built = lf.build(ENTRY, ins, CONCEPTS)
    (applied,) = built["decisions"]
    assert applied == {"fact_id": "rec:a", "status": "needs_review", "decision": "needs change",
                       "note": "long term only", "reason_text": "expert: long term only", "applies_uttarakhand": "yes",
                       "reviewed_by": "TEST SAMPLE", "reviewed_at": "2026-10-21"}
    assert sorted((s["fact_id"], s["reason"]) for s in built["stale_decisions"]) == [
        ("rec:b", "quote, values or scope changed since the review"), ("rec:gone", "fact not in this load")]
    assert lf.counts(built)["status"] == {"needs_review": 1, "extracted": 1}


def test_an_approval_never_bypasses_fail_closed():
    """An approved ControlMeasure whose ingredient is not registered serves no step."""
    from kg.query.answer import render
    from kg.query.templates import Template
    from tests.query.test_answer import ctx, row
    step = {"order": 1, "operation": "spraying", "method": "spray", "timing_raw": None, "instruction_en": "Spray.",
            "conditions": [], "applies": [{"target_id": "ai:imidacloprid", "is_ai": True, "reg_status": "unknown",
                                           "value": 100, "unit": "ml", "basis": "per_ha"}]}
    approved = row("rec:cm", "ControlMeasure", status="approved", steps=[step])
    c = {"row": approved, "unknown": {}, "explicit": 0, "variety": 0}
    served, reasons = render(c, ctx(), Template("ControlMeasure", ()), "local", None, {})
    assert served is None and reasons == ["ingredient not registered (fail closed): ai:imidacloprid=unknown"]
