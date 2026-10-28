"""The frozen gold set: data/gold/gold_set_annotations_final.xlsx -> data/gold/gold_<sheet>.jsonl.

  uv run python -m kg.eval.gold            # convert, write the jsonl files and report
  uv run python -m kg.eval.gold --check    # the committed jsonl files still equal the workbook; writes nothing

The workbook is fixed outside the repo (its Changelog sheet records every change) and never edited here.
Its sha256 is pinned: any other file is refused, so a new gold version is a deliberate change to
WORKBOOK_SHA256. Conversion changes nothing: one record per non-empty row with `_sheet`, `_row` (the Excel
row number) and every column as stored in the cell (strings unstripped, numbers as numbers, empty cells null).
Formulas are refused, because a cached result is not what the annotator typed. Problems are reported,
never fixed.

One conversion, in the output only (approved 2026-10-21): `step` is stored as text in some rows and as a
number in others; the jsonl files hold it as an integer, and every conversion is logged in the manifest.
A step that is not plain digits is refused, not guessed.
"""

import argparse
import hashlib
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from kg.registry import ROOT

GOLD_DIR = ROOT / "data" / "gold"
WORKBOOK = GOLD_DIR / "gold_set_annotations_final.xlsx"
WORKBOOK_SHA256 = "e4214ccdcac7c28bda76a250f5bf0e18e60c4dbada699338374ce522f83afa27"   # 2026-10-21: 11 rows unshifted
FACT_SHEETS = ("A", "B", "C")
CHANGELOG = "Changelog"
JSON_TYPES = (str, int, float, bool, type(None))
META = ("_sheet", "_row")              # added to each record; the Changelog has its own "sheet" column


class GoldError(Exception):
    pass


def out_name(sheet: str) -> str:
    return "gold_%s.jsonl" % sheet.lower() if sheet == CHANGELOG else "gold_%s.jsonl" % sheet


def read_workbook(path: Path = WORKBOOK, sha256: str | None = WORKBOOK_SHA256) -> dict:
    """{"sha256", "sheets": {name: {"header": [...], "records": [...], "empty_rows": [...]}}}."""
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if sha256 is not None and digest != sha256:
        raise GoldError("%s has sha256 %s, not the frozen %s; the gold set is never edited here"
                        % (path.name, digest, sha256))
    wb = load_workbook(path, data_only=False)
    names = [ws.title for ws in wb.worksheets]
    if sha256 is not None and names != [*FACT_SHEETS, CHANGELOG]:
        raise GoldError("sheets %s, expected %s" % (names, [*FACT_SHEETS, CHANGELOG]))
    sheets = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows())
        if not rows:
            raise GoldError("sheet %s is empty" % ws.title)
        header = [c.value for c in rows[0]]
        while header and header[-1] is None:                     # trailing unused columns
            header.pop()
        if not all(isinstance(h, str) and h for h in header) or len(set(header)) != len(header) \
                or set(header) & set(META):
            raise GoldError("sheet %s: header must be distinct non-empty names other than %s: %s"
                            % (ws.title, META, header))
        records, empty = [], []
        for cells in rows[1:]:
            r = cells[0].row
            for c in cells:
                if c.data_type == "f":
                    raise GoldError("%s!%s holds a formula (%s); refused" % (ws.title, c.coordinate, c.value))
                if c.value is not None and not isinstance(c.value, JSON_TYPES):
                    raise GoldError("%s!%s holds a %s, which JSON cannot keep as stored"
                                    % (ws.title, c.coordinate, type(c.value).__name__))
            beyond = [c.coordinate for c in cells[len(header):] if c.value is not None]
            if beyond:
                raise GoldError("sheet %s: values outside the header's columns at %s" % (ws.title, beyond))
            values = [c.value for c in cells[:len(header)]] + [None] * (len(header) - len(cells))
            if all(v is None for v in values):
                empty.append(r)
                continue
            records.append({"_sheet": ws.title, "_row": r, **dict(zip(header, values))})
        sheets[ws.title] = {"header": header, "records": records, "empty_rows": empty}
    return {"sha256": digest, "sheets": sheets}


def render(book: dict) -> dict[str, str]:
    """File name -> exact text of the jsonl files and the manifest (the output, steps as integers)."""
    out, log = output(book)
    files = {out_name(name): "".join(json.dumps(r, ensure_ascii=False) + "\n" for r in s["records"])
             for name, s in out["sheets"].items()}
    manifest = {"workbook": WORKBOOK.name, "sha256": book["sha256"],
                "sheets": {name: {"file": out_name(name), "rows": len(s["records"]), "header": s["header"],
                                  "empty_rows_skipped": s["empty_rows"]} for name, s in book["sheets"].items()},
                "conversions": log}
    files["gold_manifest.json"] = json.dumps(manifest, indent=2, ensure_ascii=False) + "\n"
    return files


def output(book: dict) -> tuple[dict, list[dict]]:
    """The book as written to jsonl: fact-sheet `step` text of plain digits becomes an integer. Returns the
    converted copy and the log of every conversion; the book read from the workbook is left as stored."""
    log, sheets = [], {}
    for name, s in book["sheets"].items():
        records = []
        for r in s["records"]:
            r = dict(r)
            if name in FACT_SHEETS and isinstance(r.get("step"), str):
                if not r["step"].isdigit():
                    raise GoldError("%s %s: step %r is not a whole number" % (name, r["fact_id"], r["step"]))
                log.append({"sheet": name, "row": r["_row"], "fact_id": r["fact_id"], "field": "step",
                            "from": r["step"], "to": int(r["step"])})
                r["step"] = int(r["step"])
            records.append(r)
        sheets[name] = {**s, "records": records}
    return {**book, "sheets": sheets}, log


def check(book: dict, directory: Path = GOLD_DIR) -> list[str]:
    """Files that are missing or differ from a fresh conversion."""
    return [name for name, text in render(book).items()
            if not (directory / name).exists() or (directory / name).read_text(encoding="utf-8") != text]


def problems(book: dict) -> dict:
    """What a reader of the gold set should know. Reported only; nothing is changed."""
    facts = [r for name in FACT_SHEETS for r in book["sheets"][name]["records"]]
    ids = Counter(r["fact_id"] for r in facts)
    types = Counter(r["fact_type"] for r in facts)
    no_quote = [r["fact_id"] for r in facts if r["quote"] is None]
    # an empty quote with text in other_values looks like a row whose cells sit one column to the left
    shifted = [r["fact_id"] for r in facts if r["quote"] is None and r["other_values"] is not None]
    mixed = {}
    for col in book["sheets"][FACT_SHEETS[0]]["header"]:
        kinds = Counter(type(r[col]).__name__ for r in facts if r[col] is not None)
        if len(kinds) > 1:
            mixed[col] = {k: [r["fact_id"] for r in facts if type(r[col]).__name__ == k][:8] for k in kinds}
    headers = {name: book["sheets"][name]["header"] for name in FACT_SHEETS}
    return {
        "rows": {name: len(book["sheets"][name]["records"]) for name in book["sheets"]},
        "fact_types": dict(types),
        "duplicate_fact_ids": sorted(i for i, n in ids.items() if n > 1),
        "headers_differ": len({tuple(h) for h in headers.values()}) > 1,
        "no_quote": no_quote,
        "possible_column_shift": shifted,
        "mixed_types": mixed,
        "empty_rows_skipped": {name: s["empty_rows"] for name, s in book["sheets"].items() if s["empty_rows"]},
    }


def main(argv=None) -> int:
    p = argparse.ArgumentParser(prog="python -m kg.eval.gold", description="Convert the frozen gold set to jsonl")
    p.add_argument("--check", action="store_true", help="compare the committed files; write nothing")
    a = p.parse_args(argv)
    try:
        book = read_workbook()
    except GoldError as e:
        print("REFUSED:", e)
        return 1
    if a.check:
        stale = check(book)
        print("gold jsonl %s" % ("up to date" if not stale else "differs from the workbook: %s" % stale))
        return 1 if stale else 0
    try:
        files = render(book)
        out, log = output(book)
    except GoldError as e:
        print("REFUSED:", e)
        return 1
    for name, text in files.items():
        (GOLD_DIR / name).write_text(text, encoding="utf-8", newline="\n")
    print("as stored:", json.dumps(problems(book), ensure_ascii=False))
    print("output   :", json.dumps(problems(out), ensure_ascii=False))
    print("step converted to an integer in %d rows (logged in gold_manifest.json)" % len(log))
    return 0


if __name__ == "__main__":
    sys.exit(main())
