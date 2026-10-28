"""The expert review workbook (step 16): one row per current fact, for an agronomist.

Rows are built from step 13's inputs (kg.load.facts.build, the facts the loader writes), so each row's hash
(kg.load.facts.row_hash) is the one the loader and the import check. The current status, review reasons and
Jev score come from the database; a fact that is not in it yet, or is flagged missing_in_latest_run, stops
the sheet. Only the reviewer columns are unlocked; row_hash is a hidden column, and a hidden "meta" sheet
names the source. The optional AI pre-check columns (--precheck) are a locked hint, matched by fact_id, never
read by the import.
"""

import datetime as dt
import re
from collections import Counter

from kg.load.facts import row_hash
from kg.schema.vocab import SCOPE_IN_UNMAPPED_REASON
from kg.verify.claims import window_text

ORDER = ("DryingSpec", "StorageRule", "HarvestRule", "QualitySpec", "VarietyRecommendation", "SowingWindow",
         "NutrientDose", "SeedRate", "Spacing", "SeedTreatment", "IrrigationRule", "PestThreshold", "ControlMeasure",
         "Caution")                                       # VarietyRecommendation, SowingWindow: schema v1.6
FACTS, GUIDE, META = "Facts", "How to review", "meta"
# (key, header, width); the reviewer columns are the only editable cells
COLUMNS = (("review_no", "Review no.", 8), ("fact_id", "Fact id", 22), ("subtype", "Subtype", 15),
           ("instruction", "Instruction (code-built from the stored fact)", 60), ("crop", "Crop", 10),
           ("scope", "Scope (pest / target, variety / duration, purpose, soil, region)", 34),
           ("values", "Values (with the printed text)", 50), ("quote", "Quote (exact, from the book)", 60),
           ("pages", "PDF page / printed page", 12), ("source", "Source", 24), ("status", "Current status", 14),
           ("review_reasons", "Review reasons", 50), ("jev_score", "Jev score", 9),
           ("unmapped", "Unmapped phrases", 30),
           ("ai_precheck", "AI pre-check", 11), ("ai_note", "AI note", 45),
           ("decision", "DECISION: approve / reject / needs change", 18),
           ("note", "REASON OR CORRECTION", 40),
           ("uttarakhand", "APPLIES IN UTTARAKHAND? yes / no / not sure", 18),
           ("reviewer", "REVIEWER NAME", 18), ("date", "DATE (yyyy-mm-dd)", 14),
           ("row_hash", "row_hash", 10))
REVIEWER = ("decision", "note", "uttarakhand", "reviewer", "date")
OPTIONAL = ("ai_precheck", "ai_note")    # a hint only: never imported, and a workbook without them still reads
GUIDE_LINES = (
    "You are checking whether each row says exactly what the book says, and whether it is safe to give to a farmer.",
    "Read the quote and page. If the instruction matches the book and is good advice: approve.",
    "If the book is wrong or outdated, or the advice is unsafe: reject, and say why.",
    "If it's almost right but the scope or a value needs fixing: needs change, and write the correction.",
    "\"Applies in Uttarakhand?\" is separate: this book is from Tamil Nadu. Say whether the advice also holds for "
    "Uttarakhand conditions.",
    "Don't edit any other column.",
    "If the Unmapped column shows a condition (e.g. 'for hybrids', 'in BPH prone areas'), the instruction as shown "
    "applies everywhere. Approve only if the advice is correct without that condition; otherwise choose "
    "'needs change' and write the condition.",
    "Approving a pesticide row records that it matches the book; it is still not shown to farmers until the "
    "product's official registration status has been checked.",
    "The AI pre-check column is a hint, not a decision: it was made by an AI reading the quote. Please check FIX "
    "and CHECK rows carefully, and don't rely on OK rows without reading them.",
)
UTTARAKHAND_LINE = GUIDE_LINES[4]


def guide_lines(entry, state_name: str | None = None) -> tuple:
    """The guide, with the Uttarakhand line naming where this source is from (2026-10-24; until then it said
    Tamil Nadu for every source). The TNAU guide is unchanged."""
    if entry.state_id == "state:33":
        return GUIDE_LINES
    where = "from %s" % state_name if entry.state_id else "national (all States/UTs)"
    line = ("\"Applies in Uttarakhand?\" is separate: this source is %s. Say whether the advice holds for "
            "Uttarakhand conditions." % where)
    return tuple(line if x == UTTARAKHAND_LINE else x for x in GUIDE_LINES)


SCOPE_FILL = "F4B183"            # orange: the Unmapped cell of a fact flagged "scope in unmapped"
PRECHECK = ("OK", "CHECK", "FIX")
PRECHECK_FILL = {"FIX": "F8696B", "CHECK": "FFEB84"}       # red, yellow; OK has no fill
PRECHECK_FIELDS = ["review_no", "fact_id", "ai_precheck", "ai_note"]
REVIEWER_HEAD, REVIEWER_CELL, HEAD = "FFD966", "FFF2CC", "D9E1F2"
SHEET_VERSION = "1"


class SheetError(Exception):
    def __init__(self, problems: list[str]):
        super().__init__("\n".join(problems))
        self.problems = problems


def statement(claim: str) -> str:
    """A code-built Jev claim as a plain statement: "Does the passage state that X?" -> "X."."""
    s = re.sub(r"^Does the passage state that,?\s*", "", claim.strip()).rstrip("?").strip()
    return s[:1].upper() + s[1:] + "."


def instruction(rec: dict) -> str:
    claims = rec["props"].get("verifier_claims") or []
    if claims:
        return "\n".join(statement(c) for c in claims)
    fact = rec["fact"]
    steps = [s["instruction_en"] for s in fact["steps"] if s.get("instruction_en")]
    text = fact.get("text_en") or "; ".join(steps) or fact.get("summary_en") or ""
    return "(no code-built claim; as extracted) " + text


def _name(names: dict, cid: str) -> str:
    return names.get(cid, cid)


def scope(fact: dict, names: dict) -> str:
    parts = []
    # the pests and targets the fact is for (2026-10-22): without them two IPM releases read alike
    for label, key in (("pest", "pest_ids"), ("target", "target_ids"), ("about", "about_ids")):
        if fact.get(key):
            parts.append("%s: %s" % (label, ", ".join(_name(names, x) for x in fact[key])))
    if fact["subtype"] == "ControlMeasure" and not fact.get("target_ids"):
        parts.append("target: none (general practice)")
    if fact["subtype"] == "QualitySpec":                    # schema v1.5
        parts.append("commodity: %s" % fact["commodity_form"])
        parts.append("grades: %s (%s%s)" % (", ".join(fact["grades"]), fact["grades_origin"],
                                            ": %s" % fact["grades_evidence"] if fact.get("grades_evidence") else ""))
        if fact.get("season_label"):
            parts.append("season: %s" % fact["season_label"])
    if fact["subtype"] == "SowingWindow":                   # schema v1.6
        parts.append("window for: %s" % (fact.get("window_for") or "sowing").replace("_", " "))
    for label, key in (("variety", "variety_ids"), ("season", "season_ids"), ("system", "system_ids"),
                       ("soil", "soil_ids"), ("purpose", "purpose_ids"), ("stage", "stage_ids")):
        if fact.get(key):
            parts.append("%s: %s" % (label, ", ".join(_name(names, x) for x in fact[key])))
    if fact.get("duration_group"):
        parts.append("duration: %s" % fact["duration_group"])
    parts.append("region: %s" % ", ".join(_name(names, r["id"]) for r in fact["regions"]))
    if fact.get("excluded_regions"):
        parts.append("except: %s" % ", ".join(r["raw"] for r in fact["excluded_regions"]))
    return "\n".join(parts)


TEXT_FIELDS = ("maturity_signs", "structure", "storage_term_raw", "criterion")


def values(fact: dict, names: dict) -> str:
    """Every stored value with the text the book prints for it."""
    lines = ["%s: %s (printed: %s)" % (k, fact.get(k), raw) for k, raw in sorted((fact.get("raw") or {}).items())]
    lines += ["%s: %s" % (k, "; ".join(fact[k]) if isinstance(fact[k], list) else fact[k])
              for k in TEXT_FIELDS if fact.get(k)]
    lines += ["if: %s" % c["raw"] for c in fact.get("conditions") or []]
    if fact["subtype"] == "VarietyRecommendation":          # schema v1.6: the varieties the printed names link to
        lines.append("recommends: %s" % ", ".join(_name(names, v) for v in fact["recommends"]))
    if fact["subtype"] == "SowingWindow":                   # schema v1.6: the dates read from the printed words
        lines.append("window: %s" % window_text(fact).replace("from ", "", 1))
    if fact["subtype"] == "QualitySpec":                    # schema v1.5: the limit as stored and as printed
        lines.append("%s (printed: %s): %s%s%s" % (
            fact["parameter"], fact["parameter_raw"], fact["limit_type"],
            "" if fact.get("limit_value") is None else " %g%s%s" % (
                fact["limit_value"], "%" if fact.get("unit") == "pct" else "",
                "" if fact.get("limit_inclusive", True) else " (not inclusive)"),
            ", part of %s" % fact["part_of"] if fact.get("part_of") else ""))
    for s in fact["steps"]:
        head = "step %d%s: %s" % (s["order"], " option %d" % s["alt_option"] if s.get("alt_option") else "",
                                  s.get("operation") or "")
        if s.get("method"):
            head += " (%s)" % s["method"]
        if s.get("timing_raw"):
            head += ", %s" % s["timing_raw"]
        for k in ("duration_raw", "repeat_raw"):                       # schema v1.4 duration; repeats
            if s.get(k):
                head += ", %s" % s[k]
        if s.get("purpose_raw"):
            head += " (%s)" % s["purpose_raw"]
        lines.append(head)
        for a in s["applies"]:
            extra = [x for x in (a.get("conc_raw"), a.get("ratio_raw"), a.get("share_raw"), a.get("mix_conc_raw")) if x]
            if a.get("carrier_raw"):
                extra.append("with %s" % a["carrier_raw"])
            lines.append("    %s: %s" % (_name(names, a["target_id"]), "; ".join([a.get("raw") or ""] + extra).strip("; ")))
        lines += ["    if: %s" % c["raw"] for c in s.get("conditions") or []]
    for i, c in enumerate(fact.get("clauses") or []):
        lines.append("clause %d: %s %s" % (i + 1, c.get("observation") or "", c.get("raw") or ""))
    return "\n".join(lines)


def rows(entry, built: dict, current: dict, names: dict) -> list[dict]:
    """One row per fact of the load, post-harvest first. `current` is the database's view of each fact id:
    status, review_reasons, verifier_score, missing_in_latest_run."""
    problems, out = [], []
    for rec in built["recs"]:
        db = current.get(rec["id"])
        if db is None:
            problems.append("%s (%s) is not in the database; load it first" % (rec["id"], rec["short"]))
            continue
        if db.get("missing_in_latest_run"):
            problems.append("%s is flagged missing_in_latest_run" % rec["id"])
            continue
        fact = rec["fact"]
        ev = fact["evidence"]
        out.append({
            "fact_id": rec["id"], "subtype": rec["subtype"], "instruction": instruction(rec),
            "crop": ", ".join(_name(names, c) for c in fact["crop_ids"]), "scope": scope(fact, names),
            "values": values(fact, names), "quote": "\n".join(e["quote"] for e in ev),
            "pages": "\n".join("%s / %s" % (e["pdf_page"], e.get("printed_page") or "-") for e in ev),
            "source": entry.title, "status": db["status"],
            "review_reasons": "\n".join(db.get("review_reasons") or []), "jev_score": db.get("verifier_score"),
            "unmapped": "\n".join(fact.get("unmapped") or []),
            "scope_flag": SCOPE_IN_UNMAPPED_REASON in (db.get("review_reasons") or []),
            "row_hash": row_hash(fact), "_page": ev[0]["pdf_page"] if ev else 0})
    unknown = sorted({r["subtype"] for r in out} - set(ORDER))
    if unknown:
        problems.append("no review order for subtypes %s" % unknown)
    if problems:
        raise SheetError(problems)
    out.sort(key=lambda r: (ORDER.index(r["subtype"]), r["_page"], r["fact_id"]))
    for i, r in enumerate(out, 1):
        r["review_no"] = i
    return out


def read_precheck(path) -> list[dict]:
    """The external AI pre-check notes (review_no, fact_id, ai_precheck, ai_note), as stored."""
    import csv
    with open(path, encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames != PRECHECK_FIELDS:
            raise SheetError(["%s: columns %s, expected %s" % (path, reader.fieldnames, PRECHECK_FIELDS)])
        return list(reader)


def add_precheck(sheet_rows: list[dict], notes: list[dict]) -> list[str]:
    """Put each note on its row, matched by fact_id. A row without a note keeps blank cells, and a note may have a
    blank result (an earlier note that no longer applies; 2026-10-22). Refused (SheetError, rows untouched) when
    a note's fact_id is not a row, a fact_id repeats, or a result is not OK / CHECK / FIX or blank. Returns the
    rows whose review_no differs from the note's (matching is by fact_id, so these are only reported)."""
    by_id = {r["fact_id"]: r for r in sheet_rows}
    problems, seen = [], Counter()
    for n in notes:
        seen[n["fact_id"]] += 1
        if n["fact_id"] not in by_id:
            problems.append("note for %s (review_no %s): no such fact in the sheet" % (n["fact_id"], n["review_no"]))
        if n["ai_precheck"] not in PRECHECK + ("",):
            problems.append("%s: result %r is not one of %s or blank" % (n["fact_id"], n["ai_precheck"],
                                                                         ", ".join(PRECHECK)))
        if not n["ai_precheck"] and not n["ai_note"]:
            problems.append("%s: a note with neither a result nor a text" % n["fact_id"])
    problems += ["%s: %d notes" % (i, k) for i, k in seen.items() if k > 1]
    if problems:
        raise SheetError(problems)
    moved = []
    for n in notes:
        r = by_id[n["fact_id"]]
        r["ai_precheck"], r["ai_note"] = n["ai_precheck"] or None, n["ai_note"]
        if str(r["review_no"]) != n["review_no"]:
            moved.append("%s: review_no %s in the sheet, %s in the notes" % (n["fact_id"], r["review_no"], n["review_no"]))
    return moved


def current_from_db(driver, database=None) -> dict:
    found, _, _ = driver.execute_query(
        "MATCH (r:Recommendation) RETURN r.id AS id, r.status AS status, r.review_reasons AS review_reasons, "
        "r.verifier_score AS verifier_score, coalesce(r.missing_in_latest_run, false) AS missing_in_latest_run",
        database_=database)
    return {r["id"]: r.data() for r in found}


def write_workbook(path, entry, sheet_rows: list[dict], database: str | None, state_name: str | None = None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    guide = wb.active
    guide.title = GUIDE
    guide["A1"] = "How to review"
    guide["A1"].font = Font(bold=True, size=14)
    for i, line in enumerate(guide_lines(entry, state_name), 3):
        guide.cell(i, 1, "• " + line).alignment = Alignment(wrap_text=True, vertical="top")
    guide.column_dimensions["A"].width = 120

    ws = wb.create_sheet(FACTS)
    keys = [k for k, _, _ in COLUMNS]
    col = {k: i for i, k in enumerate(keys, 1)}
    for k, header, width in COLUMNS:
        c = ws.cell(1, col[k], header)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor=REVIEWER_HEAD if k in REVIEWER else HEAD)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(col[k])].width = width
    for n, r in enumerate(sheet_rows, 2):
        for k in keys:
            c = ws.cell(n, col[k], None if k in REVIEWER else r.get(k))
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if k in REVIEWER:
                c.protection = Protection(locked=False)
                c.fill = PatternFill("solid", fgColor=REVIEWER_CELL)
            if k == "date":
                c.number_format = "yyyy-mm-dd"
        if r["scope_flag"]:
            ws.cell(n, col["unmapped"]).fill = PatternFill("solid", fgColor=SCOPE_FILL)
        if r.get("ai_precheck") in PRECHECK_FILL:
            for k in OPTIONAL:
                ws.cell(n, col[k]).fill = PatternFill("solid", fgColor=PRECHECK_FILL[r["ai_precheck"]])
    last = len(sheet_rows) + 1
    letter = {k: get_column_letter(i) for k, i in col.items()}
    for k, options in (("decision", "approve,reject,needs change"), ("uttarakhand", "yes,no,not sure")):
        v = DataValidation(type="list", formula1='"%s"' % options, allow_blank=True, showErrorMessage=True,
                           errorTitle="Not allowed", error="Choose one of: %s" % options.replace(",", ", "))
        ws.add_data_validation(v)
        v.add("%s2:%s%d" % (letter[k], letter[k], last))
    dv = DataValidation(type="date", operator="greaterThanOrEqual", formula1="DATE(2026,1,1)", allow_blank=True,
                        showErrorMessage=True, errorTitle="Not a date", error="Enter a date as yyyy-mm-dd")
    ws.add_data_validation(dv)
    dv.add("%s2:%s%d" % (letter["date"], letter["date"], last))
    ws.column_dimensions[letter["row_hash"]].hidden = True
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = "A1:%s%d" % (letter["unmapped"], last)
    ws.protection.sheet = True
    ws.protection.autoFilter = False             # filtering stays allowed on the protected sheet
    ws.protection.formatColumns = False
    ws.protection.formatRows = False

    meta = wb.create_sheet(META)
    for i, (k, v) in enumerate((("source_id", entry.id), ("sheet_version", SHEET_VERSION),
                                ("created_at", dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")),
                                ("database", database or "(home)"), ("rows", len(sheet_rows))), 1):
        meta.cell(i, 1, k)
        meta.cell(i, 2, v)
    meta.sheet_state = "hidden"
    wb.active = 0
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def read_workbook(path) -> tuple[dict, list[dict]]:
    """(meta, rows) of a filled workbook; rows carry the Excel row number and every column by key."""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    if FACTS not in wb.sheetnames or META not in wb.sheetnames:
        raise SheetError(["%s: not a review workbook (sheets %s)" % (path, wb.sheetnames)])
    meta = {r[0]: r[1] for r in wb[META].iter_rows(values_only=True) if r[0]}
    ws = wb[FACTS]
    header = [c.value for c in ws[1]]
    by_header = {h: k for k, h, _ in COLUMNS}
    missing = [h for k, h, _ in COLUMNS if h not in header and k not in OPTIONAL]
    if missing:
        raise SheetError(["%s: columns missing or renamed: %s" % (path, missing)])
    out = []
    for n, values_ in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        row = {by_header[h]: v for h, v in zip(header, values_) if h in by_header}
        if row.get("fact_id"):
            out.append({**row, "_row": n})
    return meta, out
